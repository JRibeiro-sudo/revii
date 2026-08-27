# -*- coding: utf-8 -*-
"""Part 7: fixes found by verification."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
OUT="/home/user/revii/financials/Revii_VoiceCare_Financial_Model.xlsx"
FONT="Arial"; BLUE="0000FF"; BLACK="000000"; GREEN="008000"
YEL=PatternFill("solid",fgColor="FFFF00")
NUM1='#,##0.0;(#,##0.0);"-"'; EUR='#,##0;(#,##0);"-"'; MULT='0.00"x"'
thin=Side(style="thin",color="D0D0D0"); BOX=Border(bottom=thin)
BOX2=Border(top=thin,bottom=thin,left=thin,right=thin)
wb=openpyxl.load_workbook(OUT)

def f(c,**kw):
    c.font=Font(name=FONT,bold=kw.get("bold",False),color=kw.get("color",BLACK),
                size=kw.get("size",10),italic=kw.get("italic",False))
    if kw.get("fill"): c.fill=kw["fill"]
    if kw.get("fmt"): c.number_format=kw["fmt"]
    if kw.get("align") or kw.get("wrap"):
        c.alignment=Alignment(horizontal=kw.get("align"),vertical="center",wrap_text=kw.get("wrap",False))
    if kw.get("border"): c.border=kw["border"]
    return c

# --- FIX 1: Scenarios footer pointed at Inputs!C6 (a section header) -> #VALUE
sc=wb["Scenarios"]
sc["C17"]="=Inputs!C7"; f(sc["C17"],color=GREEN,align="center",bold=True)
sc["D17"]="=INDEX($D$5:$F$5,1,Inputs!C7)"; f(sc["D17"],color=GREEN,bold=True)

# --- FIX 2: add an LTV horizon cap (row 105 was left empty by a collision)
i=wb["Inputs"]
i["B105"]="Cap on the LTV horizon (months)"; f(i["B105"],size=10,border=BOX)
i["C105"]=36; f(i["C105"],color=BLUE,fmt="0",align="center",fill=YEL,border=BOX)
i["I105"]=("At 18% annual churn the implied average customer life is 61 months. Valuing five years of "
           "an unproven subscription is not credible, so lifetime value is capped at this horizon.")
f(i["I105"],size=9,color="595959",wrap=True)
# restore the funding line lost to a row collision
i["B106"]="Funding structure: 100% founders' equity — no debt, no grant, no external round."
f(i["B106"],bold=True,size=10,color="1F3864")
i["B107"]=("Equity is injected month by month, only as much as is needed to hold the minimum cash buffer. "
           "The total injected is the investment figure on the Dashboard.")
f(i["B107"],italic=True,size=9,color="595959")

# --- FIX 3: Spain's first trading year is Year 2, not Year 1 (it launches month 17)
d=wb["Dashboard"]
for ref,txt in [("C28","Portugal Y1"),("D28","Portugal Y5"),("E28","Spain Y2"),("F28","Spain Y5")]:
    d[ref]=txt
d["E28"]="Spain Y2"
d["J28"]=("Portugal Y1 and Spain Y2 are each market's first trading year — Spain launches in month 17, "
          "so it has no Year 1.")
f(d["J28"],size=9,color="595959",wrap=True)
# repoint Spain's first-year columns from Year 1 (D:O) to Year 2 (P:AA)
for r,rng in [(29,"24:AA25"),(30,"34:AA34")]:
    src="24" if r==29 else "34"
    if r==29:
        d["E29"]="=IFERROR(SUM('Model ES'!P24:AA25)/SUM('Model ES'!P14:AA14),0)"
    else:
        d["E30"]="=IFERROR(SUM('Model ES'!P34:AA34)/SUM('Model ES'!P14:AA14),0)"

# --- FIX 4: cap LTV at the horizon, and stop reporting a negative payback period
d["B33"]="Average customer life (months, capped for LTV)"
for col in "CDEF":
    life = "Inputs!$C$43" if col in "CD" else "Inputs!$D$43"
    d[f"{col}33"]=f"=MIN({life},Inputs!$C$105)"
    f(d[f"{col}33"],fmt=NUM1,size=10,align="center",border=BOX2)
for col in "CDEF":
    d[f"{col}37"]=f'=IF({col}31<=0,"never",{col}35/{col}31)'
    f(d[f"{col}37"],fmt=NUM1,size=10,align="center",border=BOX2)
d["J36"]=("A LTV/COCA above 3.0x is the usual threshold for a healthy subscription business. "
          "A negative figure means the customer loses money over their life at these assumptions.")
f(d["J36"],size=9,color="595959",wrap=True)
d["J34"]=None

# --- FIX 5: flag the exit multiple's leverage on the headline return
d["J49"]=("The money multiple and IRR are driven almost entirely by the exit multiple on Inputs!C103. "
          "Treat them as a framing device, not a forecast.")
f(d["J49"],size=9,color="595959",wrap=True)

wb.save(OUT)
print("part 7 saved")
