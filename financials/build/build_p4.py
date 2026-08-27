# -*- coding: utf-8 -*-
"""Part 4: consolidated monthly P&L, capex, working capital, cash flow, balance sheet."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT="/home/user/revii/financials/Revii_VoiceCare_Financial_Model.xlsx"
FONT="Arial"; BLACK="000000"; GREEN="008000"
HDR=PatternFill("solid",fgColor="1F3864"); SEC=PatternFill("solid",fgColor="D9E2F3")
TOT=PatternFill("solid",fgColor="E2EFDA"); WARN=PatternFill("solid",fgColor="FCE4D6")
EUR='#,##0;(#,##0);"-"'; EUR2='#,##0.00;(#,##0.00);"-"'; PCT='0.0%;(0.0%);"-"'; NUM='#,##0;(#,##0);"-"'

N=60
def C(i): return get_column_letter(3+i)
MKT="INDEX(Scenarios!$D$15:$F$15,1,Inputs!$C$7)"   # marketing spend multiplier

wb=openpyxl.load_workbook(OUT)
ws=wb.create_sheet("P&L Monthly",5)
ws.sheet_view.showGridLines=False
ws.freeze_panes="D5"
ws.column_dimensions["A"].width=2
ws.column_dimensions["B"].width=44
ws.column_dimensions["C"].width=9
for i in range(1,N+1): ws.column_dimensions[C(i)].width=10

def st(c,**kw):
    c.font=Font(name=FONT,bold=kw.get("bold",False),color=kw.get("color",BLACK),
                size=kw.get("size",9),italic=kw.get("italic",False))
    if kw.get("fill"): c.fill=kw["fill"]
    if kw.get("fmt"): c.number_format=kw["fmt"]
    if kw.get("align"): c.alignment=Alignment(horizontal=kw["align"],vertical="center")
    return c
def put(ref,v,**kw):
    ws[ref]=v; return st(ws[ref],**kw)
def band(row,label):
    ws.cell(row,2,label)
    for c in range(2,3+N): st(ws.cell(row,c),bold=True,size=9,fill=SEC,color="1F3864")
def line(row,label,unit,fmla,fmt=EUR,*,first=None,bold=False,fill=None,color=BLACK,ind=0):
    put(f"B{row}","    "*ind+label,size=9,bold=bold,fill=fill)
    put(f"C{row}",unit,size=8,color="808080",align="center",fill=fill)
    for i in range(1,N+1):
        col=C(i); prev=C(i-1) if i>1 else None
        f=first if (i==1 and first is not None) else fmla
        put(f"{col}{row}",f.replace("{c}",col).replace("{p}",prev or col),
            fmt=fmt,size=9,bold=bold,fill=fill,color=color)

put("B2","CONSOLIDATED MONTHLY MODEL  —  PORTUGAL + SPAIN",bold=True,size=13,color="1F3864")
put("B3","Profit and loss, capital expenditure, working capital, cash flow and the founders' equity injection.",
    italic=True,size=9,color="595959")
put("B4","Month",bold=True,size=9,color="FFFFFF",fill=HDR)
put("C4","Unit",bold=True,size=8,color="FFFFFF",fill=HDR,align="center")
for i in range(1,N+1):
    put(f"{C(i)}4",i,bold=True,size=9,color="FFFFFF",fill=HDR,align="center",fmt="0")
line(5,"Model year","yr","=MIN(5,INT(({c}4-1)/12)+1)","0")

band(7,"PROFIT & LOSS  (EUR)")
line(8, "Net revenue — Portugal","EUR","='Model PT'!{c}26",color=GREEN,ind=1)
line(9, "Net revenue — Spain","EUR","='Model ES'!{c}26",color=GREEN,ind=1)
line(10,"Total net revenue","EUR","={c}8+{c}9",bold=True,fill=TOT)
line(11,"Cost of service — Portugal","EUR","=-'Model PT'!{c}34",color=GREEN,ind=1)
line(12,"Cost of service — Spain","EUR","=-'Model ES'!{c}34",color=GREEN,ind=1)
line(13,"Total cost of service","EUR","={c}11+{c}12")
line(14,"Gross profit","EUR","={c}10+{c}13",bold=True,fill=TOT)
line(15,"Gross margin","%","=IFERROR({c}14/{c}10,0)",PCT)

line(17,"Customer acquisition (COCA)","EUR","=-('Model PT'!{c}40+'Model ES'!{c}40)",ind=1)
line(18,"Brand marketing & content","EUR",f"=-INDEX(Inputs!$C$77:$G$77,1,{{c}}5)*{MKT}",ind=1)
line(19,"Personnel","EUR",
     "=-(INDEX(Inputs!$C$74:$G$74,1,{c}5)+INDEX(Inputs!$C$75:$G$75,1,{c}5)+INDEX(Inputs!$C$76:$G$76,1,{c}5)"
     "+IF({c}4>=Inputs!$C$14,INDEX(Inputs!$C$82:$G$82,1,{c}5),0))",ind=1)
line(20,"Cloud & tools","EUR","=-INDEX(Inputs!$C$78:$G$78,1,{c}5)",ind=1)
line(21,"General & administrative","EUR",
     "=-(INDEX(Inputs!$C$79:$G$79,1,{c}5)+INDEX(Inputs!$C$80:$G$80,1,{c}5)"
     "+INDEX(Inputs!$C$81:$G$81,1,{c}5)+INDEX(Inputs!$C$83:$G$83,1,{c}5))",ind=1)
line(22,"Total operating expenses","EUR","=SUM({c}17:{c}21)",bold=True)
line(23,"EBITDA","EUR","={c}14+{c}22",bold=True,fill=TOT)
line(24,"EBITDA margin","%","=IFERROR({c}23/{c}10,0)",PCT)
line(25,"Depreciation","EUR","=-{c}42")
line(26,"EBIT","EUR","={c}23+{c}25",bold=True)
line(27,"Tax loss pool — opening","EUR","={p}30",first="=0")
line(28,"Corporation tax","EUR","=-MAX(0,{c}26-{c}27)*Inputs!$C$100")
line(29,"Net income","EUR","={c}26+{c}28",bold=True,fill=TOT)
line(30,"Tax loss pool — closing","EUR","=MAX(0,{c}27-{c}26)")

band(32,"CAPITAL EXPENDITURE  (EUR)")
line(33,"Voice Care platform build","EUR",
     "=IF({c}4<=Inputs!$C$12,Inputs!$C$86/Inputs!$C$12,0)",ind=1)
line(34,"Preventive Care module","EUR","=IF({c}4=Inputs!$C$15,Inputs!$C$87,0)",ind=1)
line(35,"Spain localisation & compliance","EUR","=IF({c}4=Inputs!$C$14,Inputs!$C$88,0)",ind=1)
line(36,"Ongoing development capex","EUR",
     "=IF({c}5=1,0,(Inputs!$C$86+IF({c}4>Inputs!$C$15,Inputs!$C$87,0)"
     "+IF({c}4>Inputs!$C$14,Inputs!$C$88,0))*Inputs!$C$89/12)",ind=1)
line(37,"Total capital expenditure","EUR","=SUM({c}33:{c}36)",bold=True)
line(38,"Cumulative capital expenditure","EUR","={p}38+{c}37",first="=D37")
line(40,"Gross fixed assets","EUR","={c}38")
line(41,"Accumulated depreciation","EUR","={p}41+{c}42",first="=D42")
line(42,"Depreciation charge for the month","EUR",
     "={p}38/(Inputs!$C$90*12)",first="=0")
line(43,"Net fixed assets","EUR","={c}40-{c}41",bold=True)

band(45,"WORKING CAPITAL  (EUR)")
line(46,"Net revenue per active customer","EUR",
     "=IFERROR({c}10/('Model PT'!{c}14+'Model ES'!{c}14),0)",EUR2,ind=1)
line(47,"Annual-prepay cash collected up front","EUR",
     "=('Model PT'!{c}11+'Model ES'!{c}11)*Inputs!$C$95*{c}46*12*(1-Inputs!$C$96)",ind=1)
line(48,"Receivables","EUR","={c}10*(1-Inputs!$C$95)*Inputs!$C$93/Inputs!$C$16",ind=1)
line(49,"Payables","EUR","=(-{c}13-{c}22)*Inputs!$C$94/Inputs!$C$16",ind=1)
line(50,"Deferred revenue (annual prepay)","EUR","=MAX(0,{p}50+{c}47-{p}50/12)",first="=D47",ind=1)
line(51,"Net working capital","EUR","={c}48-{c}49-{c}50",bold=True)
line(52,"Change in net working capital","EUR","={c}51-{p}51",first="=D51")

band(54,"CASH FLOW  (EUR)")
line(55,"EBITDA","EUR","={c}23",ind=1)
line(56,"Corporation tax paid","EUR","={c}28",ind=1)
line(57,"Capital expenditure","EUR","=-{c}37",ind=1)
line(58,"Change in working capital","EUR","=-{c}52",ind=1)
line(59,"Free cash flow","EUR","=SUM({c}55:{c}58)",bold=True,fill=TOT)
line(60,"Cumulative free cash flow","EUR","={p}60+{c}59",first="=D59")
line(61,"Cash before founders' injection","EUR","={p}64+{c}59",first="=D59")
line(62,"Founders' equity injected","EUR","=MAX(0,Inputs!$C$97-{c}61)",bold=True,fill=WARN)
line(63,"Cumulative equity invested","EUR","={p}63+{c}62",bold=True,first="=D62")
line(64,"Cash — closing balance","EUR","={c}61+{c}62",bold=True,fill=TOT)

band(66,"BALANCE SHEET  (month end, EUR)")
line(67,"Cash","EUR","={c}64",ind=1)
line(68,"Receivables","EUR","={c}48",ind=1)
line(69,"Net fixed assets","EUR","={c}43",ind=1)
line(70,"Total assets","EUR","=SUM({c}67:{c}69)",bold=True)
line(71,"Payables","EUR","={c}49",ind=1)
line(72,"Deferred revenue","EUR","={c}50",ind=1)
line(73,"Paid-in equity","EUR","={c}63",ind=1)
line(74,"Retained earnings","EUR","={p}74+{c}29",first="=D29",ind=1)
line(75,"Total liabilities & equity","EUR","=SUM({c}71:{c}74)",bold=True)
line(76,"Balance check (must be nil)","EUR","={c}70-{c}75",bold=True)

wb.save(OUT)
print("part 4 saved")
