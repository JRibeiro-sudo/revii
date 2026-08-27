# -*- coding: utf-8 -*-
"""Part 3: Model PT and Model ES - the monthly market engines."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT="/home/user/revii/financials/Revii_VoiceCare_Financial_Model.xlsx"
FONT="Arial"; BLACK="000000"; GREEN="008000"
HDR=PatternFill("solid",fgColor="1F3864"); SEC=PatternFill("solid",fgColor="D9E2F3")
TOT=PatternFill("solid",fgColor="E2EFDA")
EUR='#,##0;(#,##0);"-"'; EUR2='#,##0.00;(#,##0.00);"-"'; PCT='0.0%;(0.0%);"-"'
NUM='#,##0;(#,##0);"-"'; NUM1='#,##0.0;(#,##0.0);"-"'; CMIN='0.000'
thin=Side(style="thin",color="BFBFBF")

N=60
def C(i): return get_column_letter(3+i)     # month i -> column

wb=openpyxl.load_workbook(OUT)

def build(sheetname, title, mkt, pos):
    """mkt: dict of Inputs refs specific to this market."""
    ws=wb.create_sheet(sheetname,pos)
    ws.sheet_view.showGridLines=False
    ws.freeze_panes="D5"
    ws.column_dimensions["A"].width=2
    ws.column_dimensions["B"].width=42
    ws.column_dimensions["C"].width=9
    for i in range(1,N+1): ws.column_dimensions[C(i)].width=10

    def st(c,**kw):
        c.font=Font(name=FONT,bold=kw.get("bold",False),color=kw.get("color",BLACK),
                    size=kw.get("size",9),italic=kw.get("italic",False))
        if kw.get("fill"): c.fill=kw["fill"]
        if kw.get("fmt"): c.number_format=kw["fmt"]
        if kw.get("align"): c.alignment=Alignment(horizontal=kw["align"],vertical="center")
        return c
    def put(ref,v,**kw):
        ws[ref]=v; return st(ws[ref],**kw)

    put("B2",title,bold=True,size=13,color="1F3864")
    put("B3","Monthly engine. Every driver comes from the Inputs sheet — nothing is hardcoded here.",
        italic=True,size=9,color="595959")

    def band(row,label):
        ws.cell(row,2,label)
        for c in range(2,3+N): st(ws.cell(row,c),bold=True,size=9,fill=SEC,color="1F3864")

    def line(row,label,unit,fmla,fmt,*,first=None,bold=False,fill=None,color=BLACK,ind=0):
        put(f"B{row}","   "*ind+label,size=9,bold=bold,fill=fill)
        put(f"C{row}",unit,size=8,color="808080",align="center",fill=fill)
        for i in range(1,N+1):
            col=C(i); prev=C(i-1) if i>1 else None
            f = first if (i==1 and first is not None) else fmla
            f = f.replace("{c}",col).replace("{p}",prev or col)
            put(f"{col}{row}",f,fmt=fmt,size=9,bold=bold,fill=fill,color=color)

    # --- timeline
    put("B4","Month",bold=True,size=9,color="FFFFFF",fill=HDR)
    put("C4","Unit",bold=True,size=8,color="FFFFFF",fill=HDR,align="center")
    for i in range(1,N+1):
        put(f"{C(i)}4",i,bold=True,size=9,color="FFFFFF",fill=HDR,align="center",fmt="0")
    line(5,"Model year","yr","=MIN(5,INT(({c}4-1)/12)+1)","0")
    line(6,"Months since this market launched","mo",
         f"=IF({{c}}4<Inputs!{mkt['launch']},0,{{c}}4-Inputs!{mkt['launch']}+1)","0")
    line(7,"Market year","yr","=IF({c}6=0,0,MIN(5,INT(({c}6-1)/12)+1))","0")

    band(9,"CUSTOMERS")
    line(10,"Opening active customers","cust","={p}13",NUM,first="=0")
    line(11,"New paying customers","cust",
         f"=IF({{c}}6<1,0,IF({{c}}6=1,Inputs!{mkt['adds0']},"
         f"MIN(Inputs!{mkt['cap']},{{p}}11*(1+INDEX(Inputs!{mkt['growth']},{{c}}7)*Inputs!$C$39))))",
         NUM, first=f"=IF(D6=1,Inputs!{mkt['adds0']},0)")
    line(12,"Customers lost to churn","cust",f"=-{{c}}10*Inputs!{mkt['churn_m']}",NUM)
    line(13,"Closing active customers","cust","={c}10+{c}11+{c}12",NUM,bold=True)
    line(14,"Average active customers","cust","=({c}10+{c}13)/2",NUM)
    line(15,"Free trials started","trials",f"=IFERROR({{c}}11/Inputs!{mkt['trial_c']},0)",NUM)

    band(17,"REVENUE  (EUR)")
    line(18,"Voice Care price (indexed)","EUR",
         f"=Inputs!{mkt['price_v']}*(1+Inputs!{mkt['price_inf']})^({{c}}5-1)",EUR2)
    line(19,"Preventive Care price (indexed)","EUR",
         f"=Inputs!{mkt['price_p']}*(1+Inputs!{mkt['price_inf']})^({{c}}5-1)",EUR2)
    line(20,"Preventive Care attach rate","%",
         f"=IF(OR({{c}}7=0,{{c}}4<Inputs!$C$15),0,INDEX(Inputs!{mkt['attach']},1,{{c}}7)*Inputs!$C$51)",PCT)
    line(21,"Gross billings — Voice Care","EUR","={c}14*{c}18",EUR,ind=1)
    line(22,"Gross billings — Preventive Care","EUR","={c}14*{c}20*{c}19",EUR,ind=1)
    line(23,"Total gross billings (what customers pay)","EUR","={c}21+{c}22",EUR)
    line(24,"Net revenue — Voice Care","EUR",f"={{c}}21*Inputs!{mkt['net_fact']}",EUR,ind=1)
    line(25,"Net revenue — Preventive Care","EUR",f"={{c}}22*Inputs!{mkt['net_fact']}",EUR,ind=1)
    line(26,"Total net revenue","EUR","={c}24+{c}25",EUR,bold=True,fill=TOT)

    band(28,"COST OF SERVICE  (EUR)")
    line(29,"Cost per voice minute","EUR","=INDEX(Inputs!$C$61:$G$61,1,{c}5)",CMIN)
    line(30,"AI voice — paying customers","EUR","={c}14*Inputs!$C$57*{c}29",EUR,ind=1)
    line(31,"AI voice — free trials","EUR",
         f"={{c}}15*Inputs!{mkt['trial_d']}*Inputs!$C$55*Inputs!$C$56*{{c}}29",EUR,ind=1)
    line(32,"Data & storage","EUR","={c}14*Inputs!$C$62",EUR,ind=1)
    line(33,"Preventive Care variable cost","EUR","={c}14*{c}20*Inputs!$C$52",EUR,ind=1)
    line(34,"Total cost of service","EUR","=SUM({c}30:{c}33)",EUR,bold=True)
    line(35,"Gross profit","EUR","={c}26-{c}34",EUR,bold=True,fill=TOT)
    line(36,"Gross margin","%","=IFERROR({c}35/{c}26,0)",PCT)

    band(38,"CUSTOMER ACQUISITION  (EUR)")
    line(39,"Cost per acquired customer","EUR",
         f"=IF({{c}}7=0,0,INDEX(Inputs!{mkt['coca']},1,{{c}}7)*Inputs!$C$70)",EUR2)
    line(40,"Acquisition spend","EUR","={c}11*{c}39",EUR)
    line(41,"Contribution after acquisition","EUR","={c}35-{c}40",EUR,bold=True,fill=TOT)
    return ws

PT=dict(launch="$C$13",adds0="$C$32",growth="$C$34:$C$38",cap="$C$40",churn_m="$C$42",
        trial_c="$C$45",trial_d="$C$44",price_v="$C$20",price_p="$C$21",price_inf="$C$24",
        net_fact="$C$27",attach="$C$49:$G$49",coca="$C$68:$G$68")
ES=dict(launch="$C$14",adds0="$D$32",growth="$D$34:$D$38",cap="$D$40",churn_m="$D$42",
        trial_c="$D$45",trial_d="$D$44",price_v="$D$20",price_p="$D$21",price_inf="$D$24",
        net_fact="$D$27",attach="$C$50:$G$50",coca="$C$69:$G$69")

build("Model PT","PORTUGAL — VOICE CARE MONTHLY MODEL",PT,3)
build("Model ES","SPAIN — VOICE CARE MONTHLY MODEL",ES,4)
wb.save(OUT)
print("part 3 saved")
